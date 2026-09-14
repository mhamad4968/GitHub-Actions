from pathlib import Path

import openpyxl

NAMES = [
    "塗装工事",
    "足場工事",
    "塗装及び足場工事",
    "修繕等工事",
    "塗装付帯工事",
    "軌道工事",
    "調査設計費",
    "外注試験費",
    "交通規制費",
    "追加工事①",
    "追加工事②",
    "追加工事③",
    "追加工事④",
    "追加工事⑤",
]

d = Path(r"C:\tmp\実行予算ver2")
cands = [
    x
    for x in d.glob("*.xlsx")
    if "システム工種" in x.name and not x.name.startswith("~$")
]
if not cands:
    raise SystemExit("xlsx not found")
path = cands[0]
wb = openpyxl.load_workbook(path)
ws = wb.active
name_set = set(NAMES)

hits = []
for r in range(3, ws.max_row + 1):
    system_name = ws.cell(r, 4).value
    himoku = ws.cell(r, 5).value
    work = ws.cell(r, 6).value
    if system_name in name_set and himoku == "外注費":
        nxt = ws.cell(r + 1, 6).value if r < ws.max_row else None
        if nxt == "法定福利費":
            continue
        hits.append(r)

for r in reversed(hits):
    ws.insert_rows(r + 1)
    section = ws.cell(r, 2).value or "施工"
    code = ws.cell(r, 3).value
    system_name = ws.cell(r, 4).value
    ws.cell(r + 1, 2).value = section
    ws.cell(r + 1, 3).value = code
    ws.cell(r + 1, 4).value = system_name
    ws.cell(r + 1, 5).value = "その他費用"
    ws.cell(r + 1, 6).value = "法定福利費"

out = path.with_name("マスタ整理（システム工種）-見本-20260914.xlsx")
try:
    wb.save(path)
    saved = path
except PermissionError:
    wb.save(out)
    saved = out
print("excel appended", len(hits), "rows saved", saved)
